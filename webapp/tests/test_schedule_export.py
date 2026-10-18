"""Solved-schedule .xlsx export (webapp/core/schedule_export.py + route)."""
import io
import unittest

import openpyxl

from webapp.core.schedule_export import build_schedule_workbook
from webapp.core.xlsx_style import CODE_STYLE, FLT_FILL

META = {"ward_name": "Ward 4B", "hospital": "Siriraj", "month": "7", "year": "2026"}
SET = {"weekends": [4, 5], "holidays": [], "saturdays": [4], "sundays": [5]}
DAYS = [1, 2, 3, 4, 5]
GROUPED = {
    "days": DAYS,
    "groups": [{"key": "rn", "label": "RN", "nurse_indices": [0, 1]},
               {"key": "pn_na", "label": "PN/NA", "nurse_indices": [2]}],
    "rows": [
        {"nurse_index": 0, "name": "Alice", "group": "rn", "cells": ["ช", "off", "vac", "flt:ด", "บ"]},
        {"nurse_index": 1, "name": "Bob", "group": "rn", "cells": ["ด", "ด", "off", "off", "ช"]},
        {"nurse_index": 2, "name": "Cara", "group": "pn_na", "cells": ["บ", "ช", "ช", "off", "off"]},
    ],
}
FLAT = {"days": DAYS, "rows": [
    {"nurse_index": 0, "name": "Alice", "cells": ["ช", "off", "vac", "ด", "บ"]},
    {"nurse_index": 1, "name": "Bob", "cells": ["ด", "ด", "off", "off", "ช"]},
]}


def _col_a(ws):
    return [(c.row, c.value) for c in ws["A"] if c.value is not None]


def _row_of(ws, name):
    return next(r for r, v in _col_a(ws) if v == name)


class ExportTests(unittest.TestCase):
    def test_header_is_ward_and_month_only(self):
        ws = build_schedule_workbook(GROUPED, META, SET).active
        self.assertEqual(ws["A1"].value, "Ward 4B")
        self.assertEqual(ws["A2"].value, "Siriraj  ·  July 2026")
        text = " ".join(str(c.value) for row in ws.iter_rows(max_row=3) for c in row if c.value)
        for banned in ("SETTINGS", "need", "Plan", "tier"):
            self.assertNotIn(banned, text)

    def test_sections_in_order_with_names_and_cells(self):
        ws = build_schedule_workbook(GROUPED, META, SET).active
        col = dict((v, r) for r, v in _col_a(ws))
        rn = next(r for r, v in _col_a(ws) if str(v).startswith("RN"))
        pn = next(r for r, v in _col_a(ws) if str(v).startswith("PN/NA"))
        self.assertLess(rn, col["Alice"]); self.assertLess(col["Alice"], col["Bob"])
        self.assertLess(col["Bob"], pn); self.assertLess(pn, col["Cara"])
        for row in GROUPED["rows"]:
            r = col[row["name"]]
            got = [ws.cell(r, 2 + i).value for i in range(5)]
            self.assertEqual(got, row["cells"])

    def test_ungrouped_grid_has_no_headings(self):
        ws = build_schedule_workbook(FLAT, META, SET).active
        names = [v for _, v in _col_a(ws)]
        self.assertNotIn("RN", " ".join(str(v) for v in names))
        self.assertNotIn("PN/NA", " ".join(str(v) for v in names))
        for n in ("Alice", "Bob"):
            self.assertIn(n, names)
        r = _row_of(ws, "Alice")
        self.assertEqual([ws.cell(r, 2 + i).value for i in range(5)], FLAT["rows"][0]["cells"])

    def test_tints_shared_with_roster(self):
        ws = build_schedule_workbook(GROUPED, META, SET).active
        r = _row_of(ws, "Alice")
        self.assertEqual(ws.cell(r, 4).fill.fgColor.rgb[-6:], CODE_STYLE["vac"][0])
        self.assertEqual(ws.cell(r, 5).fill.fgColor.rgb[-6:], FLT_FILL)

    def test_weekend_columns_tinted_in_header(self):
        ws = build_schedule_workbook(GROUPED, META, SET).active
        hdr = next(r for r, v in _col_a(ws) if str(v).startswith("Nurse"))
        self.assertNotEqual(ws.cell(hdr, 2).fill.fgColor.rgb, ws.cell(hdr, 5).fill.fgColor.rgb)
        self.assertEqual(ws.cell(hdr + 1, 2).value, "W")   # 1 Jul 2026 = Wednesday

    def test_totals_columns_from_fairness(self):
        from webapp.core.fairness import fairness_stats
        fair = [{"key": "", "label": "", "fair": fairness_stats(FLAT)}]
        ws = build_schedule_workbook(FLAT, META, SET, fair).active
        r = _row_of(ws, "Alice")
        self.assertEqual(ws.cell(r, 7).value, 3)    # ช, ด, บ
        self.assertEqual(ws.cell(r, 8).value, 1)    # one off (vac excluded)

    def test_missing_meta_still_builds(self):
        ws = build_schedule_workbook(FLAT, {}, {}).active
        self.assertTrue(ws["A1"].value)


"""Shared visual vocabulary for ShiftWork's Excel documents.

One definition of the palette, per-token tints and cell helpers, used by both the
roster templates (``tools/gen_templates.py``) and the solved-schedule export
(``webapp/core/schedule_export.py``) so a ``vac`` cell looks identical in both.
It lives under ``webapp/core`` so the dependency points from developer tooling
to the app, never the other way round. Pure openpyxl; no app imports.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col_letter  # noqa: F401  (re-export)

TEAL, TEAL_DARK, SLATE = "0F766E", "0B5952", "334155"
INK, BLUE, GRID = "1F2937", "1D4ED8", "CBD5E1"
LABEL_FILL, WHITE, BAND = "F1F5F9", "FFFFFF", "F8FAFC"
WEEKEND_FILL, WEEKEND_INK, WEEKEND_LETTER = "E0F2FE", TEAL_DARK, "0369A1"
LETTER_FILL, LETTER_INK = "E2E8F0", "6B7280"
FLT_FILL, REF_FILL, NEG_FILL = "FBCFE8", "FCA5A5", "FECDD3"
REST_BODY, EX_MUTED = "F0F9FF", "6B7280"
TYPE_STYLE = {"senior": ("FEF3C7", "92400E"), "new": ("DBEAFE", "1E40AF")}

# code -> (fill, font colour). Drives legend, Guide and conditional formatting.
CODE_STYLE = {
    "ช": ("FDE68A", "92400E"), "บ": ("FDBA74", "9A3412"), "ด": ("C7D2FE", "3730A3"),
    "ช/บ": ("A7F3D0", "065F46"), "ด/บ": ("DDD6FE", "5B21B6"),
    "off": ("E5E7EB", "6B7280"), "vac": ("A5F3FC", "155E75"),
}
LEGEND_TEXT = {
    "ช": "Day / เวรเช้า", "บ": "Evening / เวรบ่าย", "ด": "Night / เวรดึก",
    "ช/บ": "Day + Evening / เช้าควบบ่าย", "ด/บ": "Night + Evening / ดึกควบบ่าย",
    "off": "Day off / วันหยุด", "vac": "Vacation / ลาพักร้อน",
}

_SIDE = Side(style="thin", color=GRID)
BOX = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

LEFT1 = Alignment(horizontal="left", vertical="center", indent=1)
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")


def solid(hex_):
    return PatternFill("solid", fgColor=hex_)


def font(size, bold=False, color=INK):
    return Font(name="Arial", size=size, bold=bold, color=color)


def put(ws, ref, value=None, *, font=None, fill=None, align=None, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font is not None:
        c.font = font
    if fill is not None:
        c.fill = solid(fill)
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    return c

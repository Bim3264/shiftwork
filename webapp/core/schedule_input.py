"""ScheduleInput — the one canonical representation of a ward's solve input.

Both entry points (CSV upload and the in-browser grid) normalize into this
shape, and everything downstream — storage (JSON), the grid renderer, and the
solver adapter — reads it. Keeping a single representation is what stops the
upload path and the grid path from drifting apart.

The representation deliberately stores the *raw* schedule grid (the shift
symbols exactly as a user typed them), NOT the solver's post-processed request
structures. DataImporter samples/merges requests down to tier limits when it
builds the model; that transformation belongs to solving, not to the input we
store and let users re-edit.

Parsing of the sectioned CSV format reuses DataImporter's own section/key-value
parsers so the two never disagree about the file format.
"""

from __future__ import annotations

import csv
import io
import os
import re
import tempfile
from dataclasses import dataclass, field
from typing import Any

from constant import LocaleShift
from dataimporter import DataImporter

# Tokens a schedule cell may contain. Blank means "no request". 'off'/'mtg' are
# control tokens; the rest are the Thai shift symbols the solver understands.
_SHIFT_SYMBOLS = {
    LocaleShift.DAY.value,            # ช
    LocaleShift.EVENING.value,        # บ
    LocaleShift.NIGHT.value,          # ด
    LocaleShift.DAY_EVENING.value,    # ช/บ
    LocaleShift.NIGHT_EVENING.value,  # ด/บ
}
_CONTROL_TOKENS = {"off", "mtg", "vac"}
VALID_CELL_TOKENS = _SHIFT_SYMBOLS | _CONTROL_TOKENS | {""}

# Settings the constraint selector exposes, with their defaults. weekends are
# stored 1-based (as written in the CSV, e.g. "6 7 13 14"); everything that
# needs 0-based (the Solver constructor) converts at the boundary.
_DEFAULT_SETTINGS: dict[str, Any] = {
    "num_days": 31,
    "weekends": [6, 7, 13, 14, 20, 21, 27, 28],
    "allow_evening_to_night": False,
    "allow_night_to_day": False,
    "head_nurse_special_shift": True,
    # Partial request acceptance: 100 = every request required (default);
    # below 100 the solver may drop requests to reach feasibility, keeping as
    # many as possible. relax_days_off allows dropping requested days off too.
    "min_request_percent": 100,
    "relax_days_off": False,
    # When True (default) every requested vacation ('vac' cell) is guaranteed —
    # a hard constraint the solver can't drop. When False, vacations become
    # droppable like ordinary off-requests (subject to min_request_percent).
    "enforce_vacation": True,
    # Manually-added holiday day-numbers (1-based). Merged into `weekends` by the
    # month auto-fill; holidays behave like weekends for the solver.
    "extra_holidays": [],
    # Per-shift coverage minimums (nurses required per shift), by day band.
    # Defaults preserve the original hard-coded values.
    "coverage_weekday_day": 5,
    "coverage_weekday_evening": 3,
    "coverage_weekday_night": 3,
    "coverage_weekend_day": 4,
    "coverage_weekend_evening": 2,
    "coverage_weekend_night": 2,
}

_BOOL_SETTING_KEYS = (
    "allow_evening_to_night",
    "allow_night_to_day",
    "head_nurse_special_shift",
    "relax_days_off",
    "enforce_vacation",
)

# Integer settings clamped to >= 0 (coverage minimums).
_COVERAGE_KEYS = (
    "coverage_weekday_day", "coverage_weekday_evening", "coverage_weekday_night",
    "coverage_weekend_day", "coverage_weekend_evening", "coverage_weekend_night",
)


def _to_bool(value: Any, default: bool) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    return str(value).strip().lower() == "true"


def _parse_weekends_1based(value: str) -> list[int]:
    """Parse a '6 7 13 14' style string into a sorted list of 1-based ints."""
    if not str(value).strip():
        return list(_DEFAULT_SETTINGS["weekends"])
    parts = str(value).replace(",", " ").split()
    return sorted(int(p) for p in parts if p.strip())


def _parse_day_list(value) -> list[int]:
    """Parse a space/comma list of day numbers; empty → [] (unlike weekends)."""
    if isinstance(value, list):
        return sorted(int(x) for x in value)
    if not str(value).strip():
        return []
    return sorted(int(p) for p in str(value).replace(",", " ").split() if p.strip())


@dataclass
class Nurse:
    """One row of the grid: a name, a type, and any requested cells.

    shifts maps a 1-based day number to the raw cell token for that day. Days
    with no request are simply absent from the dict.

    active=False means the nurse is disabled: kept in the roster (name, type and
    requests preserved so they can be re-enabled) but excluded from the generated
    schedule. This is the "remove without deleting" behaviour.
    """

    name: str
    type: str = ""
    shifts: dict[int, str] = field(default_factory=dict)
    active: bool = True


@dataclass
class ScheduleInput:
    ward_meta: dict[str, str] = field(default_factory=dict)
    settings: dict[str, Any] = field(default_factory=lambda: dict(_DEFAULT_SETTINGS))
    nurses: list[Nurse] = field(default_factory=list)

    # ── Construction ─────────────────────────────────────────────────────────

    @classmethod
    def from_csv(cls, text: str) -> "ScheduleInput":
        """Parse a sectioned (or legacy) ShiftWork CSV into a ScheduleInput.

        Reuses DataImporter's section and key/value parsers so the file format
        is interpreted identically to the solver's own import path.
        """
        # _parse_sections reads a filename, so stage the text in a temp file.
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".csv", encoding="utf-8", delete=False
        ) as tmp:
            tmp.write(text)
            tmp_path = tmp.name
        try:
            sections, _has_sections = DataImporter._parse_sections(tmp_path)
        finally:
            os.unlink(tmp_path)

        ward_meta = DataImporter._parse_kv_section(sections.get("ward", []))
        raw_settings = DataImporter._parse_kv_section(sections.get("settings", []))
        settings = cls._coerce_settings(raw_settings)

        schedule_lines = sections.get("schedule", [])
        nurses = cls._parse_schedule_rows(schedule_lines, settings["num_days"])

        return cls(ward_meta=ward_meta, settings=settings, nurses=nurses)

    @classmethod
    def from_grid(
        cls,
        ward_meta: dict[str, str],
        settings: dict[str, Any],
        nurses: list[Nurse],
    ) -> "ScheduleInput":
        """Build from grid-editor form data. Coerces settings the same way the
        CSV path does so both routes produce an identical object."""
        return cls(
            ward_meta=dict(ward_meta),
            settings=cls._coerce_settings(settings),
            nurses=list(nurses),
        )

    @classmethod
    def from_xlsx(cls, data) -> "ScheduleInput":
        """Parse a filled ShiftWork Excel template (its Roster sheet) straight
        into a ScheduleInput — no CSV round-trip (option 2).

        The WARD INFO / SETTINGS cards are read by their label text (the English
        half of the bilingual "English / ไทย" labels), so inserting rows in the
        cards doesn't break parsing; the day grid is located by its row of day
        numbers. Produces the same object shape as from_csv / from_grid, so the
        rest of the pipeline can't tell which entry point was used.
        """
        try:
            import openpyxl  # lazy import: only the Excel path needs it
        except ImportError as exc:  # pragma: no cover
            raise ValueError(
                "Excel upload isn't available on this server (openpyxl missing)."
            ) from exc

        src = io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
        try:
            wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
        except Exception as exc:
            raise ValueError(
                f"That doesn't look like a readable .xlsx file ({exc})."
            ) from exc
        try:
            ws = wb["Roster"] if "Roster" in wb.sheetnames else wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()

        def gv(r, c):  # 1-based coordinate access into the materialized grid
            if 1 <= r <= len(rows):
                row = rows[r - 1]
                if 1 <= c <= len(row):
                    return row[c - 1]
            return None

        # --- WARD INFO + SETTINGS cards (label-driven) ---
        ward_meta: dict[str, str] = {}
        settings: dict[str, Any] = {}
        for r in range(1, min(len(rows), 40) + 1):
            # Left card: label col A, value col C.  Right card: label col H, value col P.
            for lc, vc in ((1, 3), (8, 16)):
                label = _cell_str(gv(r, lc)).split("/")[0].strip().lower()
                if not label:
                    continue
                val = gv(r, vc)
                if label == "ward name":
                    ward_meta["ward_name"] = _cell_str(val)
                elif label == "hospital":
                    ward_meta["hospital"] = _cell_str(val)
                elif label == "head nurse":
                    ward_meta["head_nurse"] = _cell_str(val)
                elif label == "line contact":
                    ward_meta["line"] = _cell_str(val)
                elif label == "month":
                    ward_meta["month"] = _cell_str(val)
                elif label == "year":
                    ward_meta["year"] = _cell_str(val)
                elif label == "plan":
                    ward_meta["tier"] = _cell_str(val)
                elif label == "days in month":
                    if _cell_str(val):
                        settings["num_days"] = int(float(val))
                elif label == "rest days":
                    settings["weekends"] = _parse_weekends_1based(_cell_str(val))
                elif label.startswith("allow") and "evening" in label:
                    if val is not None:
                        settings["allow_evening_to_night"] = _yesno(val)
                elif label.startswith("allow") and "day" in label:
                    if val is not None:
                        settings["allow_night_to_day"] = _yesno(val)
                elif "special shift" in label:
                    if val is not None:
                        settings["head_nurse_special_shift"] = _yesno(val)
                elif label.startswith("weekday need"):
                    d, e, n = _parse_triple(val)
                    settings.update(coverage_weekday_day=d,
                                    coverage_weekday_evening=e,
                                    coverage_weekday_night=n)
                elif label.startswith("weekend need"):
                    d, e, n = _parse_triple(val)
                    settings.update(coverage_weekend_day=d,
                                    coverage_weekend_evening=e,
                                    coverage_weekend_night=n)

        # --- Day grid: find the day-number header row, then read nurse rows ---
        header_r = None
        day_cols: dict[int, int] = {}
        for r in range(1, len(rows) + 1):
            found: dict[int, int] = {}
            row = rows[r - 1]
            for c in range(1, len(row) + 1):
                v = row[c - 1]
                if isinstance(v, bool):
                    continue
                if isinstance(v, int):
                    found[v] = c
                elif isinstance(v, str) and v.strip().isdigit():
                    found[int(v)] = c
            if found.get(1) and found.get(2) and found.get(3):
                header_r, day_cols = r, found
                break
        if header_r is None:
            raise ValueError(
                "Couldn't find the day-number header row in the Excel grid — "
                "is this a ShiftWork roster template?"
            )

        nurses: list[Nurse] = []
        for r in range(header_r + 1, len(rows) + 1):
            name = _cell_str(gv(r, 1))
            if not name:
                continue  # blank rows and the weekday-letter row carry no name
            ntype = _cell_str(gv(r, 2)).lower()
            if ntype not in ("senior", "new"):
                ntype = ""
            shifts: dict[int, str] = {}
            for day, col in day_cols.items():
                raw = gv(r, col)
                token = _normalize_token(raw) if raw is not None else ""
                if token:
                    shifts[day] = token
            nurses.append(Nurse(name=name, type=ntype, shifts=shifts))

        if "num_days" not in settings and day_cols:
            settings["num_days"] = max(day_cols)

        return cls.from_grid(ward_meta=ward_meta, settings=settings, nurses=nurses)

    # ── Coercion / parsing helpers ───────────────────────────────────────────

    @staticmethod
    def _coerce_settings(raw: dict[str, Any]) -> dict[str, Any]:
        s = dict(_DEFAULT_SETTINGS)
        if "num_days" in raw and str(raw["num_days"]).strip():
            s["num_days"] = int(raw["num_days"])
        if "weekends" in raw:
            wk = raw["weekends"]
            s["weekends"] = wk if isinstance(wk, list) else _parse_weekends_1based(wk)
        if "extra_holidays" in raw:
            s["extra_holidays"] = _parse_day_list(raw["extra_holidays"])
        for key in _BOOL_SETTING_KEYS:
            if key in raw:
                s[key] = _to_bool(raw[key], _DEFAULT_SETTINGS[key])
        if "min_request_percent" in raw and str(raw["min_request_percent"]).strip() != "":
            try:
                s["min_request_percent"] = max(0, min(100, int(raw["min_request_percent"])))
            except (TypeError, ValueError):
                pass
        for key in _COVERAGE_KEYS:
            if key in raw and str(raw[key]).strip() != "":
                try:
                    s[key] = max(0, int(raw[key]))
                except (TypeError, ValueError):
                    pass
        return s

    @staticmethod
    def _parse_schedule_rows(lines: list[str], num_days: int) -> list[Nurse]:
        text = "\n".join(lines).strip()
        if not text:
            return []
        reader = csv.reader(io.StringIO(text))
        rows = [r for r in reader if any(c.strip() for c in r)]
        if not rows:
            return []

        header = [h.strip() for h in rows[0]]
        has_type = len(header) > 1 and header[1].lower() == "type"
        # Day columns are the header cells that are digits; map header text ->
        # column index so extra/re-ordered columns are handled safely.
        day_cols = {
            int(h): idx for idx, h in enumerate(header) if h.isdigit()
        }

        nurses: list[Nurse] = []
        for row in rows[1:]:
            name = row[0].strip() if len(row) > 0 else ""
            if not name:
                continue
            ntype = row[1].strip() if has_type and len(row) > 1 else ""
            shifts: dict[int, str] = {}
            for day, col in day_cols.items():
                if col < len(row):
                    token = _normalize_token(row[col])
                    if token:
                        shifts[day] = token
            nurses.append(Nurse(name=name, type=ntype, shifts=shifts))
        return nurses

    # ── Roster ───────────────────────────────────────────────────────────────

    def active_nurses(self) -> list[Nurse]:
        """The nurses that go to the solver (disabled ones are excluded but kept
        in self.nurses). to_solver_csv() and the result-grid mapping both use
        this exact ordering."""
        return [n for n in self.nurses if n.active]

    # ── Emission ─────────────────────────────────────────────────────────────

    def to_solver_csv(self) -> str:
        """Emit a sectioned CSV that DataImporter/Solver parse back to this same
        schedule. This is what the solver adapter feeds to the engine."""
        out = io.StringIO()

        out.write("[ward]\n")
        for key, value in self.ward_meta.items():
            out.write(f"{key},{value}\n")

        out.write("\n[settings]\n")
        out.write(f"num_days,{self.settings['num_days']}\n")
        weekends = " ".join(str(d) for d in self.settings["weekends"])
        out.write(f"weekends,{weekends}\n")
        extra = " ".join(str(d) for d in self.settings.get("extra_holidays", []))
        out.write(f"extra_holidays,{extra}\n")
        for key in _BOOL_SETTING_KEYS:
            out.write(f"{key},{'true' if self.settings[key] else 'false'}\n")
        out.write(f"min_request_percent,{self.settings['min_request_percent']}\n")
        for key in _COVERAGE_KEYS:
            out.write(f"{key},{self.settings[key]}\n")

        out.write("\n[schedule]\n")
        num_days = self.settings["num_days"]
        writer = csv.writer(out, lineterminator="\n")
        writer.writerow(["Name", "Type"] + [str(d) for d in range(1, num_days + 1)])
        # Only active nurses go to the solver; disabled ones stay in self.nurses
        # but are omitted here so they don't appear in the schedule.
        for nurse in self.active_nurses():
            cells = [nurse.shifts.get(d, "") for d in range(1, num_days + 1)]
            writer.writerow([nurse.name, nurse.type] + cells)

        return out.getvalue()

    # ── JSON (for DB storage) ────────────────────────────────────────────────

    def to_dict(self) -> dict[str, Any]:
        return {
            "ward_meta": self.ward_meta,
            "settings": self.settings,
            "nurses": [
                # JSON object keys must be strings; days are re-int'd on load.
                {"name": n.name, "type": n.type, "active": n.active,
                 "shifts": {str(k): v for k, v in n.shifts.items()}}
                for n in self.nurses
            ],
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "ScheduleInput":
        nurses = [
            Nurse(
                name=n["name"],
                type=n.get("type", ""),
                shifts={int(k): v for k, v in n.get("shifts", {}).items()},
                active=n.get("active", True),  # default True for pre-existing data
            )
            for n in data.get("nurses", [])
        ]
        return cls(
            ward_meta=dict(data.get("ward_meta", {})),
            settings=cls._coerce_settings(data.get("settings", {})),
            nurses=nurses,
        )

    # ── Validation ───────────────────────────────────────────────────────────

    def validate(self) -> None:
        """Raise ValueError with a user-friendly message if the input is unusable.

        Kept intentionally lightweight — tier limits (max nurses etc.) are
        enforced by DataImporter during the solve, not duplicated here.
        """
        num_days = self.settings["num_days"]
        if num_days < 1:
            raise ValueError("num_days must be at least 1.")
        # Only active nurses are solved, so the roster checks apply to them.
        active = self.active_nurses()
        if not active:
            raise ValueError("At least one active nurse is required.")
        # The senior-nurse rule assigns a head (index 0) and deputy (index 1),
        # so the solver needs at least two active nurses when it is enabled;
        # fewer would crash the solver. Surface it here as a clear message.
        if self.settings["head_nurse_special_shift"] and len(active) < 2:
            raise ValueError(
                "The head/deputy senior-nurse rule needs at least 2 active nurses. "
                "Enable another nurse, or turn off 'head nurse special shift'."
            )
        for nurse in self.nurses:
            for day, token in nurse.shifts.items():
                if not (1 <= day <= num_days):
                    raise ValueError(
                        f"Nurse '{nurse.name}' has a request on day {day}, "
                        f"outside 1..{num_days}."
                    )
                if token not in VALID_CELL_TOKENS:
                    raise ValueError(
                        f"Nurse '{nurse.name}' day {day}: '{token}' is not a "
                        f"valid shift token."
                    )


def _normalize_token(raw: str) -> str:
    """Normalize a raw cell string to a canonical token.

    Control words are lower-cased ('OFF' -> 'off'); shift symbols are kept as-is.
    """
    token = str(raw).strip().replace("﻿", "")
    if token.lower() in _CONTROL_TOKENS:
        return token.lower()
    return token


def _cell_str(value) -> str:
    """A trimmed string for an Excel cell value; '' for blanks; ints stay clean
    (openpyxl gives whole numbers as floats, so 31.0 -> '31', not '31.0')."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _yesno(value) -> bool:
    """Interpret a Yes/No (or true/1/✓) settings cell as a boolean."""
    return _cell_str(value).lower() in ("yes", "y", "true", "1", "✓")


def _parse_triple(value) -> tuple[int, int, int]:
    """Parse a 'D / E / N' coverage cell (e.g. '5 / 3 / 3') into three ints."""
    nums = [int(p) for p in re.findall(r"\d+", _cell_str(value))]
    while len(nums) < 3:
        nums.append(0)
    return nums[0], nums[1], nums[2]

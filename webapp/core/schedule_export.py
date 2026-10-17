"""Solved schedule -> a printable, roster-styled .xlsx.

Pure derivation from ``SolveJob.result_grid`` (+ ward_meta / settings for the
header and the non-working-day tint); no solver run, works on old flat grids.

Layout (sheet "Schedule")
    rows 1-2   title bar: ward name / hospital, head nurse, month year
    per group  slate heading ("RN" / "PN/NA"), day-number + weekday header,
               one row per nurse (name, one cell per day, Shifts, Off)
    below      compact legend
A grid without ``groups`` (old or single-group job) renders as ONE section with
no heading. Visual vocabulary comes from ``webapp.core.xlsx_style`` (shared with
the roster templates).
"""
from __future__ import annotations

import calendar
import io
from typing import Optional

import openpyxl
from openpyxl.worksheet.properties import PageSetupProperties

from constant import NEGATE_KIND, OFFWARD_SEP, OffWard
from webapp.core.months import EN_MONTHS, parse_month_year
from webapp.core.xlsx_style import (
    BAND, BOX, CENTER, CODE_STYLE, FLT_FILL, INK, LEFT, LEFT1, LEGEND_TEXT,
    LETTER_FILL, LETTER_INK, NEG_FILL, REF_FILL, REST_BODY, SLATE, TEAL, TEAL_DARK,
    WEEKEND_FILL, WEEKEND_INK, WEEKEND_LETTER, WHITE, col_letter, font, put,
)

SHEET = "Schedule"
FIRST_DAY_COL = 2            # B; A is the nurse name
_LETTERS = "MTWTFSS"         # calendar.weekday(): Monday=0
_GROUP_LABELS = {"rn": "RN", "pn_na": "PN/NA"}
_FLT, _REF = f"{OffWard.FLOAT.value}{OFFWARD_SEP}", f"{OffWard.REFER.value}{OFFWARD_SEP}"
_NEG = f"{NEGATE_KIND}{OFFWARD_SEP}"


def _cell_style(token: str):
    """(fill, ink) for one grid cell token, from the roster's shared tints."""
    if token in CODE_STYLE:
        return CODE_STYLE[token]
    for prefix, fill_ in ((_FLT, FLT_FILL), (_REF, REF_FILL), (_NEG, NEG_FILL)):
        if token.startswith(prefix):
            return fill_, INK
    return None, INK


def _sections(grid: dict) -> list:
    """[(heading_or_None, rows, group_key)]. No ``groups`` -> one headingless section."""
    rows = grid.get("rows") or []
    groups = grid.get("groups") or []
    if not groups:
        return [(None, rows, "")]
    out, used = [], set()
    for g in groups:
        key = g.get("key", "")
        mine = [r for r in rows if r.get("group") == key]
        used.update(id(r) for r in mine)
        out.append((g.get("label") or _GROUP_LABELS.get(key, key), mine, key))
    rest = [r for r in rows if id(r) not in used]
    if rest:
        out.append((None, rest, ""))
    return out


def build_schedule_workbook(grid: dict, ward_meta: Optional[dict] = None,
                            settings: Optional[dict] = None,
                            fair: Optional[list] = None) -> openpyxl.Workbook:
    """``fair`` is ``app._fairness_for_job(job)`` (``[{key,label,fair}]``); when
    given, each nurse row ends with her Shifts / Off totals."""
    ward_meta, settings = ward_meta or {}, settings or {}
    days = list(grid.get("days") or [])
    n = len(days)
    my = parse_month_year(ward_meta)
    nonwork = {int(d) for k in ("weekends", "holidays", "saturdays", "sundays")
               for d in (settings.get(k) or ())}
    totals = {}
    for f in fair or []:
        if f.get("fair"):
            totals[f["key"]] = {p["nurse_index"]: p for p in f["fair"]["per_nurse"]}
    has_tot = bool(totals)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.sheet_view.showGridLines = False
    last_day = FIRST_DAY_COL + n - 1
    c_shifts, c_off = last_day + 1, last_day + 2
    last_col = c_off if has_tot else last_day
    last_letter = col_letter(last_col)

    ws.column_dimensions["A"].width = 24.0
    for ci in range(FIRST_DAY_COL, last_day + 1):
        ws.column_dimensions[col_letter(ci)].width = 4.6
    if has_tot:
        ws.column_dimensions[col_letter(c_shifts)].width = 7.0
        ws.column_dimensions[col_letter(c_off)].width = 6.0

    # Title: ward information and the month only.
    title = (ward_meta.get("ward_name") or "").strip() or "Shift schedule · ตารางเวร"
    month = f"{EN_MONTHS[my[1]]} {my[0]}" if my else ""
    bits = [b for b in ((ward_meta.get("hospital") or "").strip(),
                        (ward_meta.get("head_nurse") or "").strip(), month) if b]
    put(ws, "A1", title, font=font(20, True, WHITE), fill=TEAL, align=LEFT1)
    put(ws, "A2", "  ·  ".join(bits) or None, font=font(11, False, WHITE),
        fill=TEAL_DARK, align=LEFT1)
    ws.merge_cells(f"A1:{last_letter}1")
    ws.merge_cells(f"A2:{last_letter}2")
    ws.row_dimensions[1].height = 34.0
    ws.row_dimensions[2].height = 20.0

    r = 4
    for heading, rows, key in _sections(grid):
        if heading:
            put(ws, f"A{r}", f"{heading}  ·  {len(rows)}", font=font(11, True, WHITE),
                fill=SLATE, align=LEFT1)
            ws.merge_cells(f"A{r}:{last_letter}{r}")
            ws.row_dimensions[r].height = 22.0
            r += 1
        # Day-number + weekday-letter header.
        put(ws, f"A{r}", "Nurse · พยาบาล", font=font(10, True, WHITE), fill=TEAL,
            align=CENTER, border=BOX)
        ws.merge_cells(f"A{r}:A{r + 1}")
        for i, lbl in enumerate(days):
            col = col_letter(FIRST_DAY_COL + i)
            try:
                dnum = int(lbl)
            except (TypeError, ValueError):
                dnum = None
            rest = dnum in nonwork
            letter = (_LETTERS[calendar.weekday(my[0], my[1], dnum)]
                      if my and dnum and dnum <= calendar.monthrange(*my)[1] else "")
            put(ws, f"{col}{r}", lbl,
                font=font(10, True, WEEKEND_INK if rest else WHITE),
                fill=WEEKEND_FILL if rest else TEAL, align=CENTER, border=BOX)
            put(ws, f"{col}{r + 1}", letter,
                font=font(8, rest, WEEKEND_LETTER if rest else LETTER_INK),
                fill=WEEKEND_FILL if rest else LETTER_FILL, align=CENTER, border=BOX)
        if has_tot:
            for ci, text in ((c_shifts, "Shifts"), (c_off, "Off")):
                col = col_letter(ci)
                put(ws, f"{col}{r}", text, font=font(9, True, WHITE), fill=TEAL,
                    align=CENTER, border=BOX)
                ws.merge_cells(f"{col}{r}:{col}{r + 1}")
        ws.row_dimensions[r].height = 20.0
        ws.row_dimensions[r + 1].height = 14.0
        r += 2
        for i, row in enumerate(rows):
            bg = WHITE if i % 2 == 0 else BAND
            put(ws, f"A{r}", row.get("name", ""), font=font(10, True), fill=bg,
                align=LEFT1, border=BOX)
            cells = row.get("cells") or []
            for d in range(n):
                token = str(cells[d]) if d < len(cells) and cells[d] is not None else ""
                tint, ink = _cell_style(token)
                try:
                    rest = int(days[d]) in nonwork
                except (TypeError, ValueError):
                    rest = False
                put(ws, f"{col_letter(FIRST_DAY_COL + d)}{r}", token or None,
                    font=font(9, bool(tint) and token != "off", ink),
                    fill=tint or (REST_BODY if rest else bg), align=CENTER, border=BOX)
            if has_tot:
                p = totals.get(key, {}).get(row.get("nurse_index"))
                for ci, k in ((c_shifts, "shifts"), (c_off, "off_days")):
                    put(ws, f"{col_letter(ci)}{r}", p[k] if p else None,
                        font=font(10, True), fill=bg, align=CENTER, border=BOX)
            ws.row_dimensions[r].height = 18.0
            r += 1
        r += 1                                    # spacer between sections

    # Legend, below the grid (the header stays ward + month only).
    items = [(c, CODE_STYLE[c][0], CODE_STYLE[c][1], LEGEND_TEXT[c])
             for c in ("ช", "บ", "ด", "ช/บ", "ด/บ", "off", "vac")]
    extra = [("flt:ช", FLT_FILL, INK, "Float / ลอยไปช่วยหน่วยอื่น"),
             ("ref:ด", REF_FILL, INK, "Refer escort / ส่งต่อผู้ป่วย"),
             ("no:บ", NEG_FILL, INK, "Not this shift / ไม่เอาเวรนี้")]
    for line in (items, extra):
        for k, (code, fill_, ink, text) in enumerate(line):
            c0 = FIRST_DAY_COL + 4 * k
            put(ws, f"{col_letter(c0)}{r}", code, font=font(10, True, ink), fill=fill_,
                align=CENTER, border=BOX)
            put(ws, f"{col_letter(c0 + 1)}{r}", text, font=font(8), align=LEFT)
            ws.merge_cells(f"{col_letter(c0 + 1)}{r}:{col_letter(c0 + 3)}{r}")
        ws.row_dimensions[r].height = 16.0
        r += 1

    # Print: landscape, one page wide, names stay visible while scrolling.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.freeze_panes = "B3"
    return wb


def schedule_xlsx_bytes(*args, **kwargs) -> bytes:
    buf = io.BytesIO()
    build_schedule_workbook(*args, **kwargs).save(buf)
    return buf.getvalue()
